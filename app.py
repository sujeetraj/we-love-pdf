import base64
import hashlib
import io
import logging
import os
import re
import shutil
import tempfile
import time
import uuid
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import fitz
from flask import Flask, abort, g, jsonify, render_template, request, send_file
from werkzeug.utils import secure_filename


APP_NAME = "We Love PDF"
MAX_UPLOAD_MB = int(os.environ.get("MAX_UPLOAD_MB", "150"))
SESSION_TTL_SECONDS = int(os.environ.get("SESSION_TTL_MINUTES", "60")) * 60
MAX_PDF_PAGES = int(os.environ.get("MAX_PDF_PAGES", "300"))
MAX_OCR_PAGES = int(os.environ.get("MAX_OCR_PAGES", "25"))
MAX_OPERATION_SECONDS = int(os.environ.get("MAX_OPERATION_SECONDS", "120"))
TEMP_ROOT = Path(os.environ.get("APP_TEMP_ROOT", tempfile.gettempdir())) / "we-love-pdf"
ALLOWED_EXTENSIONS = {".pdf"}
TOOLS = {"merge", "split", "compress", "organize", "redact", "edit", "excel", "word", "ppt"}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_MB * 1024 * 1024

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
)
logger = logging.getLogger(APP_NAME)


def log_fingerprint(value: object, prefix: str) -> str:
    text = str(value or "")
    if not text:
        return f"{prefix}:empty"
    digest = hashlib.sha256(text.encode("utf-8", errors="replace")).hexdigest()[:12]
    return f"{prefix}:{digest}"


def log_session_id(session_id: object) -> str:
    return log_fingerprint(session_id, "session")


def log_filename(name: object) -> str:
    return log_fingerprint(name, "file")


def ensure_temp_root() -> None:
    TEMP_ROOT.mkdir(parents=True, exist_ok=True)


def clean_expired_sessions() -> None:
    ensure_temp_root()
    now = time.time()
    for path in TEMP_ROOT.iterdir():
        check_operation_budget()
        if path.is_dir() and now - path.stat().st_mtime > SESSION_TTL_SECONDS:
            shutil.rmtree(path, ignore_errors=True)


@app.before_request
def before_request() -> None:
    g.operation_started_at = time.monotonic()
    clean_expired_sessions()


@app.after_request
def add_security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    response.headers["Cache-Control"] = "no-store"
    csp = (
        "default-src 'self'; "
        "script-src 'self'; "
        "style-src 'self' 'unsafe-inline'; "
        "img-src 'self' data: blob:; "
        "connect-src 'self'; "
        "object-src 'none'; base-uri 'self'; frame-ancestors 'none'"
    )
    response.headers["Content-Security-Policy"] = csp
    return response


def session_dir(session_id: str) -> Path:
    if not re.fullmatch(r"[a-f0-9-]{36}", session_id or ""):
        abort(400, "Invalid session")
    ensure_temp_root()
    path = TEMP_ROOT / session_id
    path.mkdir(parents=True, exist_ok=True)
    path.touch()
    return path


def new_session_dir() -> tuple[str, Path]:
    session_id = str(uuid.uuid4())
    return session_id, session_dir(session_id)


def require_pdf(file_storage) -> None:
    name = secure_filename(file_storage.filename or "")
    if Path(name).suffix.lower() not in ALLOWED_EXTENSIONS:
        abort(400, "Only PDF files are allowed")


def check_operation_budget() -> None:
    started_at = getattr(g, "operation_started_at", None)
    if started_at is not None and time.monotonic() - started_at > MAX_OPERATION_SECONDS:
        abort(408, "PDF operation exceeded the configured time limit")


def validate_pdf_page_limit(page_count: int) -> None:
    if page_count > MAX_PDF_PAGES:
        abort(400, f"PDF has {page_count} pages; maximum allowed is {MAX_PDF_PAGES}")


def save_pdf_upload(file_storage, target_dir: Path, prefix: str = "upload") -> Path:
    require_pdf(file_storage)
    filename = secure_filename(file_storage.filename or "document.pdf")
    path = target_dir / f"{prefix}-{uuid.uuid4().hex}-{filename}"
    file_storage.save(path)
    try:
        doc = fitz.open(path)
    except Exception:
        path.unlink(missing_ok=True)
        abort(400, "Invalid PDF file")
    with doc:
        if doc.page_count < 1:
            abort(400, "Uploaded PDF has no pages")
        validate_pdf_page_limit(doc.page_count)
    return path


def output_path(target_dir: Path, name: str) -> Path:
    return target_dir / f"{uuid.uuid4().hex}-{secure_filename(name)}"


def send_pdf(path: Path, download_name: str):
    return send_download(path, download_name, "application/pdf")


def send_download(path: Path, download_name: str, mimetype: str):
    data = io.BytesIO(path.read_bytes())
    return send_file(
        data,
        as_attachment=True,
        download_name=download_name,
        mimetype=mimetype,
        max_age=0,
    )


def parse_pages(page_text: str, page_count: int) -> list[int]:
    if not page_text.strip():
        return list(range(page_count))
    pages: list[int] = []
    for part in page_text.split(","):
        token = part.strip()
        if not token:
            continue
        if "-" in token:
            start_text, end_text = token.split("-", 1)
            start = int(start_text)
            end = int(end_text)
            if start > end:
                start, end = end, start
            pages.extend(range(start, end + 1))
        else:
            pages.append(int(token))
    unique = []
    for page in pages:
        if page < 1 or page > page_count:
            abort(400, f"Page {page} is outside the document range")
        if page not in unique:
            unique.append(page)
    return [page - 1 for page in unique]


def make_zip(files: list[tuple[str, Path]]) -> io.BytesIO:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, path in files:
            archive.write(path, name)
    buffer.seek(0)
    return buffer


@app.route("/")
def index():
    return render_template("index.html", app_name=APP_NAME, max_upload_mb=MAX_UPLOAD_MB, active_tool="")


@app.route("/tool/<tool>")
def tool_page(tool: str):
    if tool not in TOOLS:
        abort(404, "Tool not found")
    return render_template("index.html", app_name=APP_NAME, max_upload_mb=MAX_UPLOAD_MB, active_tool=tool)


@app.post("/api/session")
def create_session():
    session_id, _path = new_session_dir()
    logger.info("session.created id=%s", log_session_id(session_id))
    return jsonify({"sessionId": session_id})


@app.post("/api/session/cleanup")
def cleanup_session():
    data = request.get_json(silent=True) or request.form
    session_id = data.get("sessionId", "")
    if re.fullmatch(r"[a-f0-9-]{36}", session_id or ""):
        shutil.rmtree(TEMP_ROOT / session_id, ignore_errors=True)
        logger.info("session.cleaned id=%s", log_session_id(session_id))
    return ("", 204)


@app.post("/api/merge")
def merge_pdfs():
    sid = request.form.get("sessionId", "")
    workdir = session_dir(sid)
    uploads = request.files.getlist("files")
    if len(uploads) < 2:
        abort(400, "Upload at least two PDFs")
    result = fitz.open()
    for index, upload in enumerate(uploads):
        check_operation_budget()
        path = save_pdf_upload(upload, workdir, "merge")
        with fitz.open(path) as doc:
            page_count = doc.page_count
            result.insert_pdf(doc)
        logger.info("merge.input index=%s pages=%s name=%s", index, page_count, log_filename(upload.filename))
    out = output_path(workdir, "merged.pdf")
    result.save(out, garbage=4, deflate=True, clean=True)
    result.close()
    logger.info("merge.output session=%s file=%s", log_session_id(sid), log_filename(out.name))
    return send_pdf(out, "we-love-pdf-merged.pdf")


@app.post("/api/split")
def split_pdf():
    sid = request.form.get("sessionId", "")
    workdir = session_dir(sid)
    upload = request.files.get("file")
    if upload is None:
        abort(400, "Upload a PDF")
    path = save_pdf_upload(upload, workdir, "split")
    mode = request.form.get("mode", "selected")
    ranges = request.form.get("pages", "")
    outputs: list[tuple[str, Path]] = []
    with fitz.open(path) as doc:
        if mode == "all":
            page_sets = [[i] for i in range(doc.page_count)]
        else:
            page_sets = [parse_pages(ranges, doc.page_count)]
        for set_index, pages in enumerate(page_sets, start=1):
            check_operation_budget()
            split_doc = fitz.open()
            for page in pages:
                check_operation_budget()
                split_doc.insert_pdf(doc, from_page=page, to_page=page)
            out = output_path(workdir, f"split-{set_index}.pdf")
            split_doc.save(out, garbage=4, deflate=True, clean=True)
            split_doc.close()
            outputs.append((f"split-{set_index}.pdf", out))
    if len(outputs) == 1:
        return send_pdf(outputs[0][1], outputs[0][0])
    archive = make_zip(outputs)
    return send_file(archive, as_attachment=True, download_name="we-love-pdf-split.zip", mimetype="application/zip")


@app.post("/api/compress")
def compress_pdf():
    sid = request.form.get("sessionId", "")
    workdir = session_dir(sid)
    upload = request.files.get("file")
    if upload is None:
        abort(400, "Upload a PDF")
    path = save_pdf_upload(upload, workdir, "compress")
    out = output_path(workdir, "compressed.pdf")
    level = request.form.get("level", "balanced")
    with fitz.open(path) as doc:
        save_kwargs = {"garbage": 4, "deflate": True, "clean": True}
        if level == "maximum":
            save_kwargs.update({"deflate_images": True, "deflate_fonts": True})
        doc.save(out, **save_kwargs)
    logger.info(
        "compress.output session=%s input_bytes=%s output_bytes=%s",
        log_session_id(sid),
        path.stat().st_size,
        out.stat().st_size,
    )
    return send_pdf(out, "we-love-pdf-compressed.pdf")


@app.post("/api/document")
def upload_document():
    sid = request.form.get("sessionId", "")
    workdir = session_dir(sid)
    upload = request.files.get("file")
    if upload is None:
        abort(400, "Upload a PDF")
    path = save_pdf_upload(upload, workdir, "document")
    token = path.stem
    with fitz.open(path) as doc:
        pages = [{"page": i + 1, "width": doc[i].rect.width, "height": doc[i].rect.height} for i in range(doc.page_count)]
    return jsonify({"documentId": token, "pages": pages})


def find_document(workdir: Path, document_id: str) -> Path:
    safe = secure_filename(document_id)
    matches = list(workdir.glob(f"{safe}.pdf"))
    if not matches:
        matches = [p for p in workdir.glob("document-*.pdf") if p.stem == safe]
    if not matches:
        abort(404, "Document not found")
    return matches[0]


def selected_page_refs(data: dict, workdir: Path) -> list[dict]:
    refs = data.get("pages", [])
    if not isinstance(refs, list) or not refs:
        abort(400, "Upload and keep at least one page in the preview")
    selected = []
    for ref in refs:
        check_operation_budget()
        document_id = str(ref.get("documentId", ""))
        page_number = int(ref.get("page", 0))
        path = find_document(workdir, document_id)
        with fitz.open(path) as doc:
            validate_pdf_page_limit(doc.page_count)
            if page_number < 1 or page_number > doc.page_count:
                abort(400, "Invalid page in preview")
        selected.append({"documentId": document_id, "page": page_number})
    return selected


def selected_document_refs(data: dict, workdir: Path) -> list[str]:
    refs = data.get("documents", [])
    if not isinstance(refs, list) or len(refs) < 2:
        abort(400, "Upload and keep at least two PDFs in the preview")
    selected = []
    for document_id in refs:
        check_operation_budget()
        safe_id = str(document_id)
        find_document(workdir, safe_id)
        selected.append(safe_id)
    return selected


def document_from_document_refs(workdir: Path, refs: list[str]) -> fitz.Document:
    result = fitz.open()
    for document_id in refs:
        check_operation_budget()
        path = find_document(workdir, document_id)
        with fitz.open(path) as source:
            validate_pdf_page_limit(source.page_count)
            result.insert_pdf(source)
    return result


def document_from_refs(workdir: Path, refs: list[dict]) -> fitz.Document:
    result = fitz.open()
    for ref in refs:
        check_operation_budget()
        path = find_document(workdir, ref["documentId"])
        page_index = int(ref["page"]) - 1
        with fitz.open(path) as source:
            result.insert_pdf(source, from_page=page_index, to_page=page_index)
    return result


def organized_document_from_items(workdir: Path, items: list[dict]) -> fitz.Document:
    if not items:
        abort(400, "Keep at least one page in the organizer")
    result = fitz.open()
    default_width = 595
    default_height = 842
    for item in items:
        check_operation_budget()
        item_type = str(item.get("type") or "page")
        rotation = int(item.get("rotation") or 0) % 360
        if item_type == "blank":
            width = float(item.get("width") or default_width)
            height = float(item.get("height") or default_height)
            page = result.new_page(width=width, height=height)
            if rotation:
                page.set_rotation(rotation)
            continue

        document_id = str(item.get("documentId", ""))
        page_number = int(item.get("page") or 0)
        path = find_document(workdir, document_id)
        with fitz.open(path) as source:
            validate_pdf_page_limit(source.page_count)
            if page_number < 1 or page_number > source.page_count:
                abort(400, "Invalid page in organizer")
            result.insert_pdf(source, from_page=page_number - 1, to_page=page_number - 1)
            inserted = result[-1]
            if rotation:
                inserted.set_rotation((inserted.rotation + rotation) % 360)
            default_width = source[page_number - 1].rect.width
            default_height = source[page_number - 1].rect.height
    return result


def span_styles(page: fitz.Page) -> list[dict]:
    spans = []
    for block in page.get_text("dict").get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                spans.append(
                    {
                        "rect": fitz.Rect(span.get("bbox", (0, 0, 0, 0))),
                        "style": excel_style_from_flags(int(span.get("flags", 0))),
                        "color": pdf_color_to_rgb(int(span.get("color", 0))),
                    }
                )
    return spans


def pdf_color_to_rgb(color: int) -> tuple[int, int, int]:
    return ((color >> 16) & 255, (color >> 8) & 255, color & 255)


def excel_style_from_flags(flags: int) -> int:
    bold = bool(flags & 16)
    italic = bool(flags & 2)
    if bold and italic:
        return 3
    if bold:
        return 1
    if italic:
        return 2
    return 0


def style_for_word(spans: list[dict], x: float, y: float) -> int:
    point = fitz.Point(x, y)
    for span in spans:
        if point in span["rect"]:
            return int(span["style"])
    return 0


def color_for_word(spans: list[dict], x: float, y: float) -> tuple[int, int, int]:
    point = fitz.Point(x, y)
    for span in spans:
        if point in span["rect"]:
            return span["color"]
    return (0, 0, 0)


def positioned_text_runs_from_page(page: fitz.Page) -> list[dict]:
    runs = []
    for block in page.get_text("dict").get("blocks", []):
        check_operation_budget()
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = str(span.get("text", "")).strip()
                if not text:
                    continue
                rect = fitz.Rect(span.get("bbox", (0, 0, 0, 0)))
                runs.append(
                    {
                        "text": text,
                        "rect": rect,
                        "style": excel_style_from_flags(int(span.get("flags", 0))),
                        "color": pdf_color_to_rgb(int(span.get("color", 0))),
                        "size": max(float(span.get("size", 12)), 6),
                        "font": str(span.get("font", "Calibri")),
                    }
                )
    return runs


def cluster_columns(x_positions: list[float]) -> list[float]:
    columns: list[float] = []
    for x in sorted(x_positions):
        if not columns or abs(x - columns[-1]) > 24:
            columns.append(x)
        else:
            columns[-1] = (columns[-1] + x) / 2
    return columns or [0]


def nearest_column(columns: list[float], x: float) -> int:
    return min(range(len(columns)), key=lambda index: abs(columns[index] - x)) + 1


def worksheet_cells_from_page(page: fitz.Page) -> tuple[list[list[dict]], list[float]]:
    words = sorted(page.get_text("words"), key=lambda word: (round(word[1], 1), word[0]))
    styles = span_styles(page)
    grouped_rows: list[list[dict]] = []
    current_y: float | None = None
    for word in words:
        check_operation_budget()
        x0, y0, x1, y1, text = word[:5]
        midpoint_y = (float(y0) + float(y1)) / 2
        if current_y is None or abs(midpoint_y - current_y) > 4:
            grouped_rows.append([])
            current_y = midpoint_y
        grouped_rows[-1].append(
            {
                "x0": float(x0),
                "x1": float(x1),
                "text": str(text),
                "style": style_for_word(styles, (float(x0) + float(x1)) / 2, midpoint_y),
                "color": color_for_word(styles, (float(x0) + float(x1)) / 2, midpoint_y),
            }
        )

    row_chunks: list[list[dict]] = []
    x_positions: list[float] = []
    for row in grouped_rows:
        chunks: list[dict] = []
        current_words: list[str] = []
        current_x = 0.0
        current_style = 0
        current_color = (0, 0, 0)
        previous_x1: float | None = None
        for word in sorted(row, key=lambda item: item["x0"]):
            if previous_x1 is not None and word["x0"] - previous_x1 > 18 and current_words:
                chunks.append({"x": current_x, "text": " ".join(current_words), "style": current_style, "color": current_color})
                x_positions.append(current_x)
                current_words = []
                current_style = 0
                current_color = (0, 0, 0)
            if not current_words:
                current_x = word["x0"]
                current_style = word["style"]
                current_color = word["color"]
            else:
                current_style = max(current_style, word["style"])
            current_words.append(word["text"])
            previous_x1 = word["x1"]
        if current_words:
            chunks.append({"x": current_x, "text": " ".join(current_words), "style": current_style, "color": current_color})
            x_positions.append(current_x)
        if chunks:
            row_chunks.append(chunks)

    if not row_chunks:
        return [[{"col": 1, "text": "No selectable text found on this page", "style": 0}]], [0]

    columns = cluster_columns(x_positions)
    sparse_rows: list[list[dict]] = []
    for row in row_chunks:
        sparse_rows.append(
            [
                {
                    "col": nearest_column(columns, cell["x"]),
                    "text": cell["text"],
                    "style": cell["style"],
                    "color": cell["color"],
                }
                for cell in row
            ]
        )
    return sparse_rows, columns


def extract_page_images(doc: fitz.Document, page: fitz.Page) -> list[dict]:
    images = []
    for image in page.get_images(full=True):
        check_operation_budget()
        xref = image[0]
        try:
            pix = fitz.Pixmap(doc, xref)
            if pix.n - pix.alpha > 3:
                pix = fitz.Pixmap(fitz.csRGB, pix)
            data = pix.tobytes("png")
        except Exception:
            logger.info("excel.image_skipped xref=%s", xref)
            continue
        for rect in page.get_image_rects(xref):
            images.append({"rect": rect, "bytes": data, "extension": "png"})
    return images


def column_name(index: int) -> str:
    name = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name


def worksheet_xml(rows: list[list[dict]], column_count: int, has_drawing: bool) -> str:
    cols = "".join(f'<col min="{index}" max="{index}" width="18" customWidth="1"/>' for index in range(1, column_count + 1))
    row_xml = []
    for row_index, row in enumerate(rows, start=1):
        cells = []
        for cell in row:
            ref = f"{column_name(int(cell['col']))}{row_index}"
            style = int(cell.get("style", 0))
            cells.append(f'<c r="{ref}" s="{style}" t="inlineStr"><is><t>{escape(str(cell["text"]))}</t></is></c>')
        row_xml.append(f'<row r="{row_index}">{"".join(cells)}</row>')
    drawing = '<drawing r:id="rIdDrawing1"/>' if has_drawing else ""
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f"<cols>{cols}</cols>"
        f'<sheetData>{"".join(row_xml)}</sheetData>{drawing}'
        "</worksheet>"
    )


def pdf_rect_to_excel_anchor(rect: fitz.Rect) -> dict:
    emu_per_pixel = 9525
    col_px = 96
    row_px = 20
    left = rect.x0 * 96 / 72
    top = rect.y0 * 96 / 72
    right = rect.x1 * 96 / 72
    bottom = rect.y1 * 96 / 72
    return {
        "from_col": int(left // col_px),
        "from_col_off": int((left % col_px) * emu_per_pixel),
        "from_row": int(top // row_px),
        "from_row_off": int((top % row_px) * emu_per_pixel),
        "to_col": int(right // col_px),
        "to_col_off": int((right % col_px) * emu_per_pixel),
        "to_row": int(bottom // row_px),
        "to_row_off": int((bottom % row_px) * emu_per_pixel),
    }


def drawing_xml(images: list[dict]) -> str:
    anchors = []
    for index, image in enumerate(images, start=1):
        anchor = pdf_rect_to_excel_anchor(image["rect"])
        anchors.append(
            '<xdr:twoCellAnchor editAs="oneCell">'
            f'<xdr:from><xdr:col>{anchor["from_col"]}</xdr:col><xdr:colOff>{anchor["from_col_off"]}</xdr:colOff><xdr:row>{anchor["from_row"]}</xdr:row><xdr:rowOff>{anchor["from_row_off"]}</xdr:rowOff></xdr:from>'
            f'<xdr:to><xdr:col>{anchor["to_col"]}</xdr:col><xdr:colOff>{anchor["to_col_off"]}</xdr:colOff><xdr:row>{anchor["to_row"]}</xdr:row><xdr:rowOff>{anchor["to_row_off"]}</xdr:rowOff></xdr:to>'
            '<xdr:pic>'
            f'<xdr:nvPicPr><xdr:cNvPr id="{index}" name="Image {index}"/><xdr:cNvPicPr/></xdr:nvPicPr>'
            f'<xdr:blipFill><a:blip r:embed="rId{index}"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
            '<xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>'
            '</xdr:pic><xdr:clientData/></xdr:twoCellAnchor>'
        )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'{"".join(anchors)}</xdr:wsDr>'
    )


def drawing_rels_xml(images: list[dict], first_media_index: int) -> str:
    rels = []
    for index, image in enumerate(images, start=1):
        media_index = first_media_index + index - 1
        rels.append(
            f'<Relationship Id="rId{index}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image{media_index}.{image["extension"]}"/>'
        )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'{"".join(rels)}</Relationships>'
    )


def worksheet_rels_xml(drawing_index: int) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rIdDrawing1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing{drawing_index}.xml"/>'
        "</Relationships>"
    )


def original_upload_stem(path: Path) -> str:
    parts = path.name.split("-", 2)
    original = parts[2] if len(parts) == 3 else path.name
    return secure_filename(Path(original).stem or "document")


def create_excel_from_refs(workdir: Path, refs: list[dict]) -> tuple[Path, str]:
    first_path = find_document(workdir, str(refs[0].get("documentId", "")))
    original_stem = original_upload_stem(first_path)
    out = output_path(workdir, f"converted-{original_stem}.xlsx")
    sheets: list[dict] = []
    for index, ref in enumerate(refs, start=1):
        check_operation_budget()
        document_id = str(ref.get("documentId", ""))
        page_number = int(ref.get("page", 0))
        path = find_document(workdir, document_id)
        with fitz.open(path) as doc:
            validate_pdf_page_limit(doc.page_count)
            if page_number < 1 or page_number > doc.page_count:
                abort(400, "Invalid page in preview")
            page = doc[page_number - 1]
            cells, columns = worksheet_cells_from_page(page)
            sheets.append(
                {
                    "name": f"Page {index}",
                    "cells": cells,
                    "column_count": max(len(columns), 1),
                    "images": extract_page_images(doc, page),
                }
            )

    sheet_defs = "".join(
        f'<sheet name="{escape(sheet["name"])}" sheetId="{index}" r:id="rId{index}"/>'
        for index, sheet in enumerate(sheets, start=1)
    )
    workbook_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f"<sheets>{sheet_defs}</sheets>"
        "</workbook>"
    )
    rels_xml = "".join(
        f'<Relationship Id="rId{index}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{index}.xml"/>'
        for index in range(1, len(sheets) + 1)
    )
    rels_xml += '<Relationship Id="rIdStyles" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
    workbook_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f"{rels_xml}</Relationships>"
    )
    sheet_overrides = "".join(
        f'<Override PartName="/xl/worksheets/sheet{index}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        for index in range(1, len(sheets) + 1)
    )
    drawing_overrides = "".join(
        f'<Override PartName="/xl/drawings/drawing{index}.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
        for index, sheet in enumerate(sheets, start=1)
        if sheet["images"]
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Default Extension="png" ContentType="image/png"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        f"{sheet_overrides}{drawing_overrides}</Types>"
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        "</Relationships>"
    )
    styles_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<fonts count="4">'
        '<font><sz val="11"/><name val="Calibri"/></font>'
        '<font><b/><sz val="11"/><name val="Calibri"/></font>'
        '<font><i/><sz val="11"/><name val="Calibri"/></font>'
        '<font><b/><i/><sz val="11"/><name val="Calibri"/></font>'
        '</fonts>'
        '<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
        '<borders count="1"><border/></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        '<cellXfs count="4">'
        '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        '<xf numFmtId="0" fontId="1" fillId="0" borderId="0" xfId="0" applyFont="1"/>'
        '<xf numFmtId="0" fontId="2" fillId="0" borderId="0" xfId="0" applyFont="1"/>'
        '<xf numFmtId="0" fontId="3" fillId="0" borderId="0" xfId="0" applyFont="1"/>'
        '</cellXfs>'
        "</styleSheet>"
    )

    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("xl/workbook.xml", workbook_xml)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        archive.writestr("xl/styles.xml", styles_xml)
        media_index = 1
        for index, sheet in enumerate(sheets, start=1):
            images = sheet["images"]
            archive.writestr(
                f"xl/worksheets/sheet{index}.xml",
                worksheet_xml(sheet["cells"], int(sheet["column_count"]), bool(images)),
            )
            if images:
                archive.writestr(f"xl/worksheets/_rels/sheet{index}.xml.rels", worksheet_rels_xml(index))
                archive.writestr(f"xl/drawings/drawing{index}.xml", drawing_xml(images))
                archive.writestr(f"xl/drawings/_rels/drawing{index}.xml.rels", drawing_rels_xml(images, media_index))
                for image in images:
                    archive.writestr(f"xl/media/image{media_index}.{image['extension']}", image["bytes"])
                    media_index += 1
    return out, f"converted-{original_stem}.xlsx"


def selected_page_packages(workdir: Path, refs: list[dict]) -> tuple[str, list[dict]]:
    first_path = find_document(workdir, str(refs[0].get("documentId", "")))
    original_stem = original_upload_stem(first_path)
    pages = []
    for index, ref in enumerate(refs, start=1):
        check_operation_budget()
        document_id = str(ref.get("documentId", ""))
        page_number = int(ref.get("page", 0))
        path = find_document(workdir, document_id)
        with fitz.open(path) as doc:
            validate_pdf_page_limit(doc.page_count)
            if page_number < 1 or page_number > doc.page_count:
                abort(400, "Invalid page in preview")
            page = doc[page_number - 1]
            cells, columns = worksheet_cells_from_page(page)
            pages.append(
                {
                    "name": f"Page {index}",
                    "width": float(page.rect.width),
                    "height": float(page.rect.height),
                    "cells": cells,
                    "text_runs": positioned_text_runs_from_page(page),
                    "column_count": max(len(columns), 1),
                    "images": extract_page_images(doc, page),
                }
            )
    return original_stem, pages


def word_run_xml(text: str, style: int) -> str:
    props = ""
    if style in {1, 3}:
        props += "<w:b/>"
    if style in {2, 3}:
        props += "<w:i/>"
    if props:
        props = f"<w:rPr>{props}</w:rPr>"
    return f'<w:r>{props}<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'


def word_image_xml(rel_id: str, rect: fitz.Rect) -> str:
    width = max(int(rect.width * 12700), 1)
    height = max(int(rect.height * 12700), 1)
    return (
        "<w:p><w:r><w:drawing><wp:inline distT=\"0\" distB=\"0\" distL=\"0\" distR=\"0\" "
        "xmlns:wp=\"http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing\">"
        f'<wp:extent cx="{width}" cy="{height}"/>'
        '<wp:docPr id="1" name="PDF image"/>'
        '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:nvPicPr><pic:cNvPr id="0" name="PDF image"/><pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip r:embed="{rel_id}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        '<pic:spPr><a:xfrm><a:off x="0" y="0"/>'
        f'<a:ext cx="{width}" cy="{height}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
        "</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>"
    )


def word_table_xml(rows: list[list[dict]]) -> str:
    table_rows = []
    for row in rows:
        cells = []
        sorted_cells = sorted(row, key=lambda item: int(item["col"]))
        for cell in sorted_cells:
            cells.append(f'<w:tc><w:p>{word_run_xml(str(cell["text"]), int(cell.get("style", 0)))}</w:p></w:tc>')
        table_rows.append(f'<w:tr>{"".join(cells)}</w:tr>')
    return (
        '<w:tbl><w:tblPr><w:tblBorders>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        f'</w:tblBorders></w:tblPr>{"".join(table_rows)}</w:tbl>'
    )


def create_word_from_refs(workdir: Path, refs: list[dict]) -> tuple[Path, str]:
    original_stem, pages = selected_page_packages(workdir, refs)
    out = output_path(workdir, f"converted-{original_stem}.docx")
    body_parts = []
    rels = []
    media_index = 1
    for page_index, page in enumerate(pages, start=1):
        if page_index > 1:
            body_parts.append('<w:p><w:r><w:br w:type="page"/></w:r></w:p>')
        body_parts.append(f'<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>{escape(page["name"])}</w:t></w:r></w:p>')
        body_parts.append(word_table_xml(page["cells"]))
        for image in page["images"]:
            rel_id = f"rIdImage{media_index}"
            body_parts.append(word_image_xml(rel_id, image["rect"]))
            rels.append(
                f'<Relationship Id="{rel_id}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image{media_index}.{image["extension"]}"/>'
            )
            media_index += 1

    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<w:body>{"".join(body_parts)}<w:sectPr><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="720" w:right="720" w:bottom="720" w:left="720"/></w:sectPr></w:body>'
        "</w:document>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Default Extension="png" ContentType="image/png"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )
    document_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'{"".join(rels)}</Relationships>'
    )
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/_rels/document.xml.rels", document_rels)
        media_index = 1
        for page in pages:
            for image in page["images"]:
                archive.writestr(f"word/media/image{media_index}.{image['extension']}", image["bytes"])
                media_index += 1
    return out, f"converted-{original_stem}.docx"


def ppt_text_body_xml(text: str, style: int) -> str:
    attrs = ""
    if style in {1, 3}:
        attrs += ' b="1"'
    if style in {2, 3}:
        attrs += ' i="1"'
    return (
        "<p:txBody><a:bodyPr/><a:lstStyle/><a:p>"
        f'<a:r><a:rPr lang="en-US" sz="1400"{attrs}/><a:t>{escape(text)}</a:t></a:r>'
        "</a:p></p:txBody>"
    )


def ppt_text_shape_xml(shape_id: int, cell: dict, row_index: int, page: dict) -> str:
    col = int(cell["col"]) - 1
    x = int((col * 140 + 28) * 12700)
    y = int((row_index * 24 + 36) * 12700)
    width = int(130 * 12700)
    height = int(22 * 12700)
    return (
        "<p:sp>"
        f'<p:nvSpPr><p:cNvPr id="{shape_id}" name="PDF text"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{width}" cy="{height}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></p:spPr>'
        f'{ppt_text_body_xml(str(cell["text"]), int(cell.get("style", 0)))}'
        "</p:sp>"
    )


def ppt_image_shape_xml(shape_id: int, rel_id: str, image: dict) -> str:
    rect = image["rect"]
    return (
        "<p:pic>"
        f'<p:nvPicPr><p:cNvPr id="{shape_id}" name="PDF image"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
        f'<p:blipFill><a:blip r:embed="{rel_id}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
        f'<p:spPr><a:xfrm><a:off x="{int(rect.x0 * 12700)}" y="{int(rect.y0 * 12700)}"/><a:ext cx="{max(int(rect.width * 12700), 1)}" cy="{max(int(rect.height * 12700), 1)}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
        "</p:pic>"
    )


def ppt_slide_xml(page: dict, slide_index: int) -> str:
    shapes = [
        '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
    ]
    shape_id = 2
    for row_index, row in enumerate(page["cells"], start=1):
        for cell in sorted(row, key=lambda item: int(item["col"])):
            shapes.append(ppt_text_shape_xml(shape_id, cell, row_index, page))
            shape_id += 1
    for image_index, image in enumerate(page["images"], start=1):
        shapes.append(ppt_image_shape_xml(shape_id, f"rIdImage{image_index}", image))
        shape_id += 1
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        f'<p:cSld name="Page {slide_index}"><p:spTree>{"".join(shapes)}</p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>'
        "</p:sld>"
    )


def ppt_slide_rels_xml(images: list[dict], first_media_index: int) -> str:
    rels = [
        '<Relationship Id="rIdLayout" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml"/>'
    ]
    for index, image in enumerate(images, start=1):
        media_index = first_media_index + index - 1
        rels.append(
            f'<Relationship Id="rIdImage{index}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image{media_index}.{image["extension"]}"/>'
        )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'{"".join(rels)}</Relationships>'
    )


def ppt_slide_master_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:sldMaster xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<p:cSld><p:spTree>'
        '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
        '</p:spTree></p:cSld>'
        '<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>'
        '<p:sldLayoutIdLst><p:sldLayoutId id="2147483649" r:id="rIdLayout1"/></p:sldLayoutIdLst>'
        '<p:txStyles><p:titleStyle/><p:bodyStyle/><p:otherStyle/></p:txStyles>'
        "</p:sldMaster>"
    )


def ppt_slide_master_rels_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdLayout1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml"/>'
        '<Relationship Id="rIdTheme1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme" Target="../theme/theme1.xml"/>'
        "</Relationships>"
    )


def ppt_slide_layout_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:sldLayout xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" type="blank" preserve="1">'
        '<p:cSld name="Blank"><p:spTree>'
        '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
        '</p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>'
        "</p:sldLayout>"
    )


def ppt_slide_layout_rels_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdMaster1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster" Target="../slideMasters/slideMaster1.xml"/>'
        "</Relationships>"
    )


def ppt_theme_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="We Love PDF">'
        '<a:themeElements><a:clrScheme name="Office">'
        '<a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1><a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>'
        '<a:dk2><a:srgbClr val="1F1F1F"/></a:dk2><a:lt2><a:srgbClr val="F7F7F8"/></a:lt2>'
        '<a:accent1><a:srgbClr val="E5322D"/></a:accent1><a:accent2><a:srgbClr val="2458D3"/></a:accent2>'
        '<a:accent3><a:srgbClr val="626262"/></a:accent3><a:accent4><a:srgbClr val="FFFFFF"/></a:accent4>'
        '<a:accent5><a:srgbClr val="B91D18"/></a:accent5><a:accent6><a:srgbClr val="171717"/></a:accent6>'
        '<a:hlink><a:srgbClr val="2458D3"/></a:hlink><a:folHlink><a:srgbClr val="5A3E8C"/></a:folHlink>'
        '</a:clrScheme><a:fontScheme name="Office"><a:majorFont><a:latin typeface="Calibri"/></a:majorFont><a:minorFont><a:latin typeface="Calibri"/></a:minorFont></a:fontScheme><a:fmtScheme name="Office"><a:fillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:fillStyleLst><a:lnStyleLst><a:ln w="6350"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln></a:lnStyleLst><a:effectStyleLst><a:effectStyle><a:effectLst/></a:effectStyle></a:effectStyleLst><a:bgFillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:bgFillStyleLst></a:fmtScheme></a:themeElements>'
        '<a:objectDefaults/><a:extraClrSchemeLst/>'
        "</a:theme>"
    )


def create_ppt_from_refs(workdir: Path, refs: list[dict]) -> tuple[Path, str]:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    original_stem, pages = selected_page_packages(workdir, refs)
    out = output_path(workdir, f"converted-{original_stem}.pptx")
    deck = Presentation()
    deck.slide_width = Emu(10058400)
    deck.slide_height = Emu(7543800)
    blank_layout = deck.slide_layouts[6]
    if len(deck.slides) == 1 and not deck.slides[0].shapes:
        deck.slides._sldIdLst.remove(deck.slides._sldIdLst[0])

    for page in pages:
        slide = deck.slides.add_slide(blank_layout)
        scale = min(float(deck.slide_width) / page["width"], float(deck.slide_height) / page["height"])

        for image in page["images"]:
            rect = image["rect"]
            slide.shapes.add_picture(
                io.BytesIO(image["bytes"]),
                Emu(int(rect.x0 * scale)),
                Emu(int(rect.y0 * scale)),
                width=Emu(max(int(rect.width * scale), 1)),
                height=Emu(max(int(rect.height * scale), 1)),
            )

        text_runs = page["text_runs"]
        if not text_runs:
            for row_index, row in enumerate(page["cells"], start=1):
                for cell in sorted(row, key=lambda item: int(item["col"])):
                    text_runs.append(
                        {
                            "text": str(cell["text"]),
                            "rect": fitz.Rect((int(cell["col"]) - 1) * 140 + 24, row_index * 24 + 32, int(cell["col"]) * 140 + 152, row_index * 24 + 54),
                            "style": int(cell.get("style", 0)),
                            "color": cell.get("color", (0, 0, 0)),
                            "size": 12,
                            "font": "Calibri",
                        }
                    )

        for item in text_runs:
            rect = item["rect"]
            box = slide.shapes.add_textbox(
                Emu(int(rect.x0 * scale)),
                Emu(int(rect.y0 * scale)),
                Emu(max(int(rect.width * scale), 1)),
                Emu(max(int(rect.height * scale), 1)),
            )
            frame = box.text_frame
            frame.margin_left = 0
            frame.margin_right = 0
            frame.margin_top = 0
            frame.margin_bottom = 0
            frame.word_wrap = False
            paragraph = frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = str(item["text"])
            run.font.size = Pt(max(float(item.get("size", 12)) * scale / 12700, 6))
            run.font.name = "Calibri"
            red, green, blue = item.get("color", (0, 0, 0))
            run.font.color.rgb = RGBColor(red, green, blue)
            style = int(item.get("style", 0))
            run.font.bold = style in {1, 3}
            run.font.italic = style in {2, 3}

    deck.save(out)
    return out, f"converted-{original_stem}.pptx"


def document_from_range(workdir: Path, document_id: str, start: int, end: int) -> fitz.Document:
    path = find_document(workdir, document_id)
    result = fitz.open()
    with fitz.open(path) as source:
        validate_pdf_page_limit(source.page_count)
        if start < 1 or end < 1 or start > source.page_count or end > source.page_count:
            abort(400, "Split range is outside the document")
        if start > end:
            start, end = end, start
        result.insert_pdf(source, from_page=start - 1, to_page=end - 1)
    return result


def search_document_text(path: Path, term: str, use_ocr: bool = False) -> tuple[list[dict], str]:
    matches: list[dict] = []
    ocr_warning = ""
    ocr_page_count = 0
    with fitz.open(path) as doc:
        validate_pdf_page_limit(doc.page_count)
        for page_index, page in enumerate(doc):
            check_operation_budget()
            rects = page.search_for(term)
            used_ocr = False
            if not rects and use_ocr and hasattr(page, "get_textpage_ocr"):
                ocr_page_count += 1
                if ocr_page_count > MAX_OCR_PAGES:
                    ocr_warning = f"OCR limited to {MAX_OCR_PAGES} pages. Narrow the document or search without OCR."
                    break
                try:
                    check_operation_budget()
                    textpage = page.get_textpage_ocr(language="eng", full=True)
                    rects = page.search_for(term, textpage=textpage)
                    used_ocr = bool(rects)
                except Exception:
                    logger.info("ocr.unavailable file=%s page=%s", log_filename(path.name), page_index + 1)
                    ocr_warning = "OCR is unavailable for this document. Search selectable text or try again without OCR."
            for rect in rects:
                matches.append(
                    {
                        "page": page_index + 1,
                        "term": term,
                        "ocr": used_ocr,
                        "rect": {
                            "x0": rect.x0,
                            "y0": rect.y0,
                            "x1": rect.x1,
                            "y1": rect.y1,
                        },
                        "pageSize": {
                            "width": page.rect.width,
                            "height": page.rect.height,
                        },
                    }
                )
    return matches, ocr_warning


def save_generated_pdf(doc: fitz.Document, workdir: Path, filename: str, maximum: bool = False) -> Path:
    out = output_path(workdir, filename)
    save_kwargs = {"garbage": 4, "deflate": True, "clean": True}
    if maximum:
        save_kwargs.update({"deflate_images": True, "deflate_fonts": True})
    check_operation_budget()
    doc.save(out, **save_kwargs)
    doc.close()
    return out


def image_bytes_from_data_url(value: str) -> bytes:
    if not value or "," not in value:
        abort(400, "Upload an image before saving image edits")
    header, payload = value.split(",", 1)
    if not header.startswith("data:image/"):
        abort(400, "Only PNG and JPEG images can be inserted")
    try:
        return base64.b64decode(payload, validate=True)
    except Exception:
        abort(400, "Invalid image data")


def hex_to_rgb(value: str, fallback: tuple[float, float, float]) -> tuple[float, float, float]:
    if not isinstance(value, str) or not re.fullmatch(r"#[0-9a-fA-F]{6}", value):
        return fallback
    return tuple(int(value[index : index + 2], 16) / 255 for index in (1, 3, 5))


def font_name(style: dict) -> str:
    family = str(style.get("fontFamily") or "helvetica")
    bold = bool(style.get("bold"))
    italic = bool(style.get("italic"))
    if family == "times":
        return "tibi" if bold and italic else "tibo" if bold else "tiit" if italic else "tiro"
    if family == "courier":
        return "cobi" if bold and italic else "cobo" if bold else "coit" if italic else "cour"
    return "hebi" if bold and italic else "hebo" if bold else "heit" if italic else "helv"


def apply_edit_action(page, action: str, text: str, rect: fitz.Rect, image_data: str = "", style: dict | None = None) -> None:
    style = style if isinstance(style, dict) else {}
    color = hex_to_rgb(str(style.get("color") or ""), (0.82, 0.08, 0.12))
    background = str(style.get("backgroundColor") or "")
    fontsize = max(6, min(float(style.get("fontSize") or 14), 96))
    if background:
        page.draw_rect(rect, color=None, fill=hex_to_rgb(background, (1, 1, 1)), overlay=True)
    if action == "text":
        if not text:
            abort(400, "Text is required")
        page.insert_textbox(rect, text, fontsize=fontsize, fontname=font_name(style), color=color, align=fitz.TEXT_ALIGN_LEFT)
        if style.get("underline"):
            underline_y = min(rect.y1 - 2, rect.y0 + fontsize + 3)
            page.draw_line(fitz.Point(rect.x0, underline_y), fitz.Point(rect.x1, underline_y), color=color, width=1)
    elif action == "highlight":
        page.add_highlight_annot(rect)
    elif action == "comment":
        page.add_text_annot(fitz.Point(rect.x0, rect.y0), text or "Comment")
    elif action == "rectangle":
        page.draw_rect(rect, color=color, width=2)
    elif action == "circle":
        page.draw_oval(rect, color=color, width=2)
    elif action == "image":
        page.insert_image(rect, stream=image_bytes_from_data_url(image_data), keep_proportion=True)
    else:
        abort(400, "Unsupported edit action")


@app.get("/api/document/<session_id>/<document_id>/thumbnail/<int:page_number>")
def thumbnail(session_id: str, document_id: str, page_number: int):
    workdir = session_dir(session_id)
    path = find_document(workdir, document_id)
    with fitz.open(path) as doc:
        validate_pdf_page_limit(doc.page_count)
        if page_number < 1 or page_number > doc.page_count:
            abort(404)
        page = doc[page_number - 1]
        check_operation_budget()
        pix = page.get_pixmap(matrix=fitz.Matrix(0.24, 0.24), alpha=False)
        return send_file(io.BytesIO(pix.tobytes("png")), mimetype="image/png", max_age=0)


@app.get("/api/document/<session_id>/<document_id>/preview/<int:page_number>")
def preview(session_id: str, document_id: str, page_number: int):
    workdir = session_dir(session_id)
    path = find_document(workdir, document_id)
    with fitz.open(path) as doc:
        validate_pdf_page_limit(doc.page_count)
        if page_number < 1 or page_number > doc.page_count:
            abort(404)
        page = doc[page_number - 1]
        check_operation_budget()
        pix = page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4), alpha=False)
        return send_file(io.BytesIO(pix.tobytes("png")), mimetype="image/png", max_age=0)


@app.post("/api/redact/search")
def search_redactions():
    data = request.get_json(force=True)
    sid = data.get("sessionId", "")
    workdir = session_dir(sid)
    document_id = str(data.get("documentId", ""))
    term = str(data.get("term", "")).strip()
    if not term:
        abort(400, "Enter text to search")
    path = find_document(workdir, document_id)
    matches, warning = search_document_text(path, term, bool(data.get("ocr")))
    return jsonify({"matches": matches, "warning": warning})


@app.post("/api/organize")
def organize_pdf():
    data = request.get_json(force=True)
    sid = data.get("sessionId", "")
    workdir = session_dir(sid)
    path = find_document(workdir, data.get("documentId", ""))
    order = data.get("order", [])
    if not order:
        abort(400, "Choose at least one page")
    with fitz.open(path) as doc:
        result = fitz.open()
        for page_number in order:
            page_index = int(page_number) - 1
            if page_index < 0 or page_index >= doc.page_count:
                abort(400, "Invalid page order")
            result.insert_pdf(doc, from_page=page_index, to_page=page_index)
        out = output_path(workdir, "organized.pdf")
        result.save(out, garbage=4, deflate=True, clean=True)
        result.close()
    return send_pdf(out, "we-love-pdf-organized.pdf")


@app.post("/api/generate/<tool>")
def generate_from_preview(tool: str):
    if tool not in TOOLS:
        abort(404, "Tool not found")
    data = request.get_json(force=True)
    sid = data.get("sessionId", "")
    workdir = session_dir(sid)
    options = data.get("options", {}) if isinstance(data.get("options", {}), dict) else {}

    if tool == "merge":
        document_refs = selected_document_refs(data, workdir)
        result = document_from_document_refs(workdir, document_refs)
        out = save_generated_pdf(result, workdir, "merged.pdf")
        logger.info("generate.output tool=%s session=%s documents=%s", tool, log_session_id(sid), len(document_refs))
        return send_pdf(out, "we-love-pdf-merged.pdf")

    if tool == "split":
        split_mode = str(options.get("splitMode") or "pages")
        if split_mode == "range":
            document_id = str(options.get("documentId") or "")
            ranges = options.get("ranges", [])
            if not document_id or not isinstance(ranges, list) or not ranges:
                abort(400, "Add at least one split range")
            if options.get("mergeRanges"):
                merged = fitz.open()
                for item in ranges:
                    check_operation_budget()
                    part = document_from_range(workdir, document_id, int(item.get("from", 1)), int(item.get("to", 1)))
                    merged.insert_pdf(part)
                    part.close()
                out = save_generated_pdf(merged, workdir, "split-ranges.pdf")
                return send_pdf(out, "we-love-pdf-split-ranges.pdf")

            outputs: list[tuple[str, Path]] = []
            for index, item in enumerate(ranges, start=1):
                check_operation_budget()
                split_doc = document_from_range(workdir, document_id, int(item.get("from", 1)), int(item.get("to", 1)))
                out = save_generated_pdf(split_doc, workdir, f"split-range-{index}.pdf")
                outputs.append((f"split-range-{index}.pdf", out))
            if len(outputs) == 1:
                return send_pdf(outputs[0][1], outputs[0][0])
            archive = make_zip(outputs)
            return send_file(archive, as_attachment=True, download_name="we-love-pdf-split-ranges.zip", mimetype="application/zip")

        refs = selected_page_refs(data, workdir)
        outputs: list[tuple[str, Path]] = []
        for index, ref in enumerate(refs, start=1):
            check_operation_budget()
            split_doc = document_from_refs(workdir, [ref])
            out = save_generated_pdf(split_doc, workdir, f"split-page-{index}.pdf")
            outputs.append((f"split-page-{index}.pdf", out))
        if len(outputs) == 1:
            return send_pdf(outputs[0][1], outputs[0][0])
        archive = make_zip(outputs)
        return send_file(archive, as_attachment=True, download_name="we-love-pdf-split.zip", mimetype="application/zip")

    if tool == "organize":
        items = options.get("items", [])
        if not isinstance(items, list):
            abort(400, "Invalid organizer items")
        result = organized_document_from_items(workdir, items)
        out = save_generated_pdf(result, workdir, "organized.pdf")
        logger.info("generate.output tool=%s session=%s items=%s", tool, log_session_id(sid), len(items))
        return send_pdf(out, "we-love-pdf-organized.pdf")

    refs = selected_page_refs(data, workdir)

    if tool == "excel":
        out, download_name = create_excel_from_refs(workdir, refs)
        logger.info("generate.output tool=%s session=%s pages=%s", tool, log_session_id(sid), len(refs))
        return send_download(
            out,
            download_name,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    if tool == "word":
        out, download_name = create_word_from_refs(workdir, refs)
        logger.info("generate.output tool=%s session=%s pages=%s", tool, log_session_id(sid), len(refs))
        return send_download(
            out,
            download_name,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    if tool == "ppt":
        out, download_name = create_ppt_from_refs(workdir, refs)
        logger.info("generate.output tool=%s session=%s pages=%s", tool, log_session_id(sid), len(refs))
        return send_download(
            out,
            download_name,
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

    result = document_from_refs(workdir, refs)

    if tool == "redact":
        rect_redactions = options.get("rects", [])
        if isinstance(rect_redactions, list):
            for item in rect_redactions:
                check_operation_budget()
                page_index = int(item.get("page") or 1) - 1
                if page_index < 0 or page_index >= result.page_count:
                    abort(400, "Redaction page is outside the preview range")
                rect = item.get("rect", {}) if isinstance(item.get("rect", {}), dict) else {}
                result[page_index].add_redact_annot(
                    fitz.Rect(
                        float(rect.get("x0") or 0),
                        float(rect.get("y0") or 0),
                        float(rect.get("x1") or 0),
                        float(rect.get("y1") or 0),
                    ),
                    fill=(0, 0, 0),
                )
        terms = [term.strip() for term in str(options.get("terms", "")).splitlines() if term.strip()]
        for page in result:
            check_operation_budget()
            for term in terms:
                for rect in page.search_for(term):
                    page.add_redact_annot(rect, fill=(0, 0, 0))
        area = options.get("area") if isinstance(options.get("area"), dict) else {}
        if area and area.get("page"):
            page_index = int(area.get("page")) - 1
            if page_index < 0 or page_index >= result.page_count:
                abort(400, "Area page is outside the preview range")
            x = float(area.get("x") or 0)
            y = float(area.get("y") or 0)
            width = float(area.get("width") or 0)
            height = float(area.get("height") or 0)
            if width <= 0 or height <= 0:
                abort(400, "Area width and height are required")
            result[page_index].add_redact_annot(fitz.Rect(x, y, x + width, y + height), fill=(0, 0, 0))
        for page in result:
            check_operation_budget()
            if page.first_annot:
                page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS)

    if tool == "edit":
        edits = options.get("edits", [])
        if isinstance(edits, list) and edits:
            for edit in edits:
                check_operation_budget()
                page_index = int(edit.get("page") or 1) - 1
                if page_index < 0 or page_index >= result.page_count:
                    abort(400, "Edit page is outside the preview range")
                x = float(edit.get("x") or 72)
                y = float(edit.get("y") or 72)
                width = float(edit.get("width") or 180)
                height = float(edit.get("height") or 40)
                rect = fitz.Rect(x, y, x + width, y + height)
                apply_edit_action(
                    result[page_index],
                    str(edit.get("action") or "text"),
                    str(edit.get("text") or "").strip(),
                    rect,
                    str(edit.get("imageData") or ""),
                    edit.get("format") if isinstance(edit.get("format"), dict) else {},
                )
        else:
            page_index = int(options.get("page") or 1) - 1
            if page_index < 0 or page_index >= result.page_count:
                abort(400, "Edit page is outside the preview range")
            x = float(options.get("x") or 72)
            y = float(options.get("y") or 72)
            width = float(options.get("width") or 180)
            height = float(options.get("height") or 40)
            rect = fitz.Rect(x, y, x + width, y + height)
            apply_edit_action(
                result[page_index],
                str(options.get("action") or "text"),
                str(options.get("text") or "").strip(),
                rect,
                str(options.get("imageData") or ""),
                options.get("format") if isinstance(options.get("format"), dict) else {},
            )

    filename = {
        "merge": "merged.pdf",
        "compress": "compressed.pdf",
        "organize": "organized.pdf",
        "redact": "redacted.pdf",
        "edit": "edited.pdf",
    }.get(tool, "output.pdf")
    out = save_generated_pdf(result, workdir, filename, maximum=options.get("level") == "maximum")
    logger.info("generate.output tool=%s session=%s pages=%s", tool, log_session_id(sid), len(refs))
    return send_pdf(out, f"we-love-pdf-{filename}")


@app.post("/api/redact")
def redact_pdf():
    sid = request.form.get("sessionId", "")
    workdir = session_dir(sid)
    upload = request.files.get("file")
    if upload is None:
        abort(400, "Upload a PDF")
    path = save_pdf_upload(upload, workdir, "redact")
    search_terms = [term.strip() for term in request.form.get("terms", "").splitlines() if term.strip()]
    page_text = request.form.get("pages", "")
    area_page = request.form.get("areaPage", "").strip()
    out = output_path(workdir, "redacted.pdf")
    redactions = 0
    with fitz.open(path) as doc:
        target_pages = parse_pages(page_text, doc.page_count) if page_text.strip() else list(range(doc.page_count))
        for page_index in target_pages:
            check_operation_budget()
            page = doc[page_index]
            for term in search_terms:
                for rect in page.search_for(term):
                    page.add_redact_annot(rect, fill=(0, 0, 0))
                    redactions += 1
        if area_page:
            page_index = int(area_page) - 1
            if page_index < 0 or page_index >= doc.page_count:
                abort(400, "Area page is outside the document range")
            x = float(request.form.get("x") or 0)
            y = float(request.form.get("y") or 0)
            width = float(request.form.get("width") or 0)
            height = float(request.form.get("height") or 0)
            if width <= 0 or height <= 0:
                abort(400, "Area width and height are required")
            doc[page_index].add_redact_annot(fitz.Rect(x, y, x + width, y + height), fill=(0, 0, 0))
            redactions += 1
        for page in doc:
            check_operation_budget()
            if page.first_annot:
                page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS)
        doc.save(out, garbage=4, deflate=True, clean=True)
    logger.info("redact.output session=%s matches=%s", log_session_id(sid), redactions)
    return send_pdf(out, "we-love-pdf-redacted.pdf")


@app.post("/api/edit")
def edit_pdf():
    sid = request.form.get("sessionId", "")
    workdir = session_dir(sid)
    upload = request.files.get("file")
    if upload is None:
        abort(400, "Upload a PDF")
    path = save_pdf_upload(upload, workdir, "edit")
    page_number = int(request.form.get("page", "1"))
    action = request.form.get("action", "text")
    text = request.form.get("text", "").strip()
    x = float(request.form.get("x", "72"))
    y = float(request.form.get("y", "72"))
    width = float(request.form.get("width", "180"))
    height = float(request.form.get("height", "40"))
    out = output_path(workdir, "edited.pdf")
    with fitz.open(path) as doc:
        if page_number < 1 or page_number > doc.page_count:
            abort(400, "Invalid page")
        page = doc[page_number - 1]
        rect = fitz.Rect(x, y, x + width, y + height)
        apply_edit_action(page, action, text, rect, request.form.get("imageData", ""))
        doc.save(out, garbage=4, deflate=True, clean=True)
    return send_pdf(out, "we-love-pdf-edited.pdf")


@app.errorhandler(400)
@app.errorhandler(404)
@app.errorhandler(408)
@app.errorhandler(413)
@app.errorhandler(500)
def handle_error(error):
    code = getattr(error, "code", 500)
    message = getattr(error, "description", "Something went wrong")
    logger.warning("request.error code=%s message=%s", code, message)
    return jsonify({"error": message}), code


if __name__ == "__main__":
    ensure_temp_root()
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "8000")), debug=False)
